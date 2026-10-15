"""AR06 · Anexos A–F del informe — DICOSUR (Taller 10, SI-886).

Genera en docs/evidencias/S10/:
  anexo_A_modelo_arquitectura.archimate   copia del modelo de Archi (Restringido)
  anexo_B_vista_proceso_aplicacion.png    vista V02
  anexo_C_estructura_informacion.xlsx     entidades, cooperación e indicadores con fórmulas
  anexo_D_arquitectura_objetivo.png       vista V04
  anexo_E_analisis_brechas.xlsx           brechas, paquetes y resumen con fórmulas
  anexo_F_seccion_5.pdf                   Sección 5 del PETI (Word o LibreOffice para exportar)

Uso (desde la raíz del repositorio):
    python 05_arquitectura/AR06_genera_anexos.py
"""
import re
import shutil
import subprocess
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
RAIZ = BASE.parent
S10 = RAIZ / "docs" / "evidencias" / "S10"
GRAFICOS = RAIZ / "graficos"
S10.mkdir(parents=True, exist_ok=True)

AZUL = "1F2D5C"
CAB = Font(bold=True, color="FFFFFF")
FONDO = PatternFill("solid", fgColor=AZUL)
BORDE = Border(*(Side(style="thin", color="999999"),) * 4)
AJUSTE = Alignment(wrap_text=True, vertical="top")
ROJO = PatternFill("solid", fgColor="F4C7C3")


def hoja(ws, df, anchos, titulo):
    ws.append([titulo])
    ws["A1"].font = Font(bold=True, size=13, color=AZUL)
    ws.append(["DICOSUR · PETI v0.10 · Restringido · Taller 10 SI-886"])
    ws["A2"].font = Font(italic=True, color="666666")
    ws.append([])
    ws.append(list(df.columns))
    for c in ws[4]:
        c.font, c.fill, c.border, c.alignment = CAB, FONDO, BORDE, AJUSTE
    for fila in df.itertuples(index=False):
        ws.append(list(fila))
        for c in ws[ws.max_row]:
            c.border, c.alignment = BORDE, AJUSTE
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"


def copia(origen, nombre):
    shutil.copyfile(origen, S10 / nombre)
    print(f"  docs/evidencias/S10/{nombre}")


def anexo_c():
    ent = pd.read_csv(BASE / "AR02_entidades.csv")
    pro = pd.read_csv(BASE / "AR02_procesos_aplicacion.csv")
    wb = Workbook()
    ws = wb.active
    ws.title = "Entidades"
    hoja(ws, ent, [20, 30, 22, 18, 24, 40, 14, 18, 34, 24], "Anexo C · Estructura de la información (AR02_entidades.csv)")
    n = 4 + len(ent)
    for fila in ws.iter_rows(min_row=5, max_row=n):
        if fila[3].value == "Ninguna":
            fila[3].fill = ROJO
        if fila[4].value == "Ninguno":
            fila[4].fill = ROJO
    ws2 = wb.create_sheet("Cooperación")
    hoja(ws2, pro, [9, 30, 38, 18, 12, 60], "Vista de cooperación proceso-aplicación (AR02_procesos_aplicacion.csv)")
    m = 4 + len(pro)
    ws3 = wb.create_sheet("Indicadores")
    ws3.append(["Indicadores de la arquitectura actual"])
    ws3["A1"].font = Font(bold=True, size=13, color=AZUL)
    ws3.append(["Calculados con fórmulas sobre las hojas Entidades y Cooperación"])
    ws3.append([])
    ws3.append(["Indicador", "N", "M"])
    for c in ws3[4]:
        c.font, c.fill, c.border = CAB, FONDO, BORDE
    filas = [
        ("Entidades sin fuente autoritativa", f'=COUNTIF(Entidades!D5:D{n},"Ninguna")', f"=COUNTA(Entidades!A5:A{n})"),
        ("Entidades sin dueño de dato", f'=COUNTIF(Entidades!E5:E{n},"Ninguno")', f"=COUNTA(Entidades!A5:A{n})"),
        ("Entidades sin fuente ni dueño", f'=COUNTIFS(Entidades!D5:D{n},"Ninguna",Entidades!E5:E{n},"Ninguno")', f"=COUNTA(Entidades!A5:A{n})"),
        ("Entidades con datos personales", f'=COUNTIF(Entidades!H5:H{n},"Sí*")', f"=COUNTA(Entidades!A5:A{n})"),
        ("Procesos con soporte deficiente", f'=COUNTIF(Cooperación!E5:E{m},"Deficiente")', f"=COUNTA(Cooperación!A5:A{m})"),
        ("Procesos con soporte parcial", f'=COUNTIF(Cooperación!E5:E{m},"Parcial")', f"=COUNTA(Cooperación!A5:A{m})"),
    ]
    for f in filas:
        ws3.append(list(f))
        for c in ws3[ws3.max_row]:
            c.border = BORDE
    ws3.column_dimensions["A"].width = 38
    wb.save(S10 / "anexo_C_estructura_informacion.xlsx")
    print("  docs/evidencias/S10/anexo_C_estructura_informacion.xlsx")


def anexo_e():
    b = pd.read_csv(BASE / "AR03_brechas.csv", dtype=str).fillna("—")
    b["esfuerzo_sem"] = b["esfuerzo_sem"].astype(int)
    p = pd.read_csv(BASE / "AR03_paquetes.csv", dtype=str).fillna("—")
    wb = Workbook()
    ws = wb.active
    ws.title = "Brechas"
    hoja(ws, b, [7, 12, 44, 22, 20, 11, 10, 16, 9, 34], "Anexo E · Análisis de brechas (AR03_brechas.csv)")
    n = 4 + len(b)
    ws2 = wb.create_sheet("Paquetes")
    p = p[["ola", "paquete", "nombre", "brechas", "depende_de", "proyecto"]]
    hoja(ws2, p, [6, 9, 40, 26, 18, 32], "Paquetes de trabajo y secuencia obligada (salida de AR03_brechas.py)")
    ws2["G4"] = "esfuerzo (sem-p)"
    ws2["G4"].font, ws2["G4"].fill, ws2["G4"].border = CAB, FONDO, BORDE
    for r in range(5, 5 + len(p)):
        ws2[f"G{r}"] = f"=SUMIF(Brechas!$I$5:$I${n},B{r},Brechas!$G$5:$G${n})"
        ws2[f"G{r}"].border = BORDE
    ws2.column_dimensions["G"].width = 14
    ws3 = wb.create_sheet("Resumen")
    ws3.append(["Brechas por capa y tipo (fórmulas)"])
    ws3["A1"].font = Font(bold=True, size=13, color=AZUL)
    ws3.append([])
    ws3.append([])
    ws3.append(["Capa", "Nuevo", "Modificar", "Retirar", "Total", "Esfuerzo (sem-p)"])
    for c in ws3[4]:
        c.font, c.fill, c.border = CAB, FONDO, BORDE
    for k, capa in enumerate(["Negocio", "Datos", "Aplicación", "Tecnología"], start=5):
        ws3.append([capa] + [f'=COUNTIFS(Brechas!$B$5:$B${n},$A{k},Brechas!$F$5:$F${n},"{t}")' for t in ("Nuevo", "Modificar", "Retirar")]
                   + [f"=SUM(B{k}:D{k})", f"=SUMIF(Brechas!$B$5:$B${n},$A{k},Brechas!$G$5:$G${n})"])
    ws3.append(["Total"] + [f"=SUM({c}5:{c}8)" for c in "BCDEF"])
    for fila in ws3.iter_rows(min_row=5, max_row=9):
        for c in fila:
            c.border = BORDE
    ws3["A9"].font = Font(bold=True)
    ws3.column_dimensions["A"].width = 16
    ws3.column_dimensions["F"].width = 16
    wb.save(S10 / "anexo_E_analisis_brechas.xlsx")
    print("  docs/evidencias/S10/anexo_E_analisis_brechas.xlsx")


# ----------------------------------------------------------------- Word / PDF
def sombrear(celda, color):
    tcpr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcpr.append(shd)


def texto_rico(par, texto):
    texto = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", texto)
    for i, trozo in enumerate(re.split(r"\*\*", texto)):
        for j, sub in enumerate(re.split(r"`", trozo)):
            run = par.add_run(sub)
            run.bold = i % 2 == 1
            if j % 2 == 1:
                run.font.name = "Consolas"


def md_a_docx(doc, md):
    lineas = md.splitlines()
    i = 0
    while i < len(lineas):
        l = lineas[i].rstrip()
        if l.startswith("# "):
            doc.add_heading(l[2:], level=1)
        elif l.startswith("## "):
            doc.add_heading(l[3:], level=2)
        elif l.startswith("### "):
            doc.add_heading(l[4:], level=3)
        elif l.startswith("|"):
            filas = []
            while i < len(lineas) and lineas[i].startswith("|"):
                if not re.match(r"^\|[-| ]+\|$", lineas[i].strip()):
                    filas.append([c.strip() for c in lineas[i].strip().strip("|").split("|")])
                i += 1
            t = doc.add_table(rows=len(filas), cols=len(filas[0]))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = False
            ncol = len(filas[0])
            minimo = [0.17 * max(len(w) for f in filas for w in (f[k].split() or [""])) + 0.35 for k in range(ncol)]
            pref = [max(len(f[k]) for f in filas) ** 0.75 for k in range(ncol)]
            libre = max(17.4 - sum(minimo), 0)
            anchos = [Cm(m + libre * x / sum(pref)) for m, x in zip(minimo, pref)]
            for a, fila in enumerate(filas):
                for b, val in enumerate(fila):
                    c = t.cell(a, b)
                    c.width = anchos[b]
                    c.text = ""
                    texto_rico(c.paragraphs[0], val)
                    for r in c.paragraphs[0].runs:
                        r.font.size = Pt(7.5)
                        if a == 0:
                            r.bold = True
                            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    if a == 0:
                        sombrear(c, AZUL)
            doc.add_paragraph()
            continue
        elif l.startswith("- "):
            texto_rico(doc.add_paragraph(style="List Bullet"), l[2:])
        elif re.match(r"^`graficos/(V0\d[^`]+\.png)`", l.strip()):
            png = re.match(r"^`graficos/(V0\d[^`]+\.png)`", l.strip()).group(1)
            doc.add_picture(str(GRAFICOS / png), width=Cm(17))
            texto_rico(doc.add_paragraph(), l)
        elif l.strip():
            texto_rico(doc.add_paragraph(), l)
        i += 1


def exporta_pdf(docx, pdf):
    """Word en Windows; LibreOffice en cualquier otro sistema."""
    ps = (f"$w=New-Object -ComObject Word.Application;$w.Visible=$false;"
          f"$d=$w.Documents.Open('{docx}');$d.ExportAsFixedFormat('{pdf}',17);$d.Close($false);$w.Quit()")
    intentos = [["powershell", "-NoProfile", "-Command", ps],
                ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(pdf.parent), str(docx)]]
    for cmd in intentos:
        try:
            subprocess.run(cmd, check=True, capture_output=True, timeout=240)
            if pdf.exists():
                return True
        except Exception:
            continue
    return False


def anexo_f():
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(1.8)
        s.top_margin = s.bottom_margin = Cm(1.8)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(9.5)
    t = doc.add_paragraph()
    r = t.add_run("Anexo F · PETI DICOSUR v0.10 · Sección 5 Arquitectura empresarial")
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor(0x1F, 0x2D, 0x5C)
    doc.add_paragraph("Taller 10 · SI-886 Planeamiento Estratégico de TI · Universidad Privada de Tacna · Documento restringido")
    md_a_docx(doc, (BASE / "5_arquitectura_empresarial.md").read_text(encoding="utf-8"))
    docx = S10 / "anexo_F_seccion_5.docx"
    pdf = S10 / "anexo_F_seccion_5.pdf"
    doc.save(docx)
    if exporta_pdf(docx, pdf):
        docx.unlink()
        print("  docs/evidencias/S10/anexo_F_seccion_5.pdf")
    else:
        print(f"  [AVISO] No se pudo exportar a PDF; queda {docx.name}")


if __name__ == "__main__":
    print("AR06 · Anexos generados:")
    copia(BASE / "DICOSUR_arquitectura.archimate", "anexo_A_modelo_arquitectura.archimate")
    copia(GRAFICOS / "V02_cooperacion_proceso_aplicacion.png", "anexo_B_vista_proceso_aplicacion.png")
    anexo_c()
    copia(GRAFICOS / "V04_arquitectura_objetivo_capas.png", "anexo_D_arquitectura_objetivo.png")
    anexo_e()
    anexo_f()
